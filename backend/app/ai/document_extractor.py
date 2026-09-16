"""
Deterministic Document Extraction Module for Pharmaceutical Complaints.

Supported Formats:
- PDF (.pdf): pypdf text extraction (detects scanned/image-only PDFs gracefully)
- DOCX (.docx): python-docx text & table extraction
- TXT (.txt): UTF-8 decoding with Latin-1 fallback
- EML (.eml): Python standard library email parser (header and body extraction)

Constraints & Rules:
- Maximum upload size: 10 MB.
- Zero production OCR: Scanned image-only PDFs return controlled, descriptive errors.
- No permanent storage: Operates strictly in memory; bytes are discarded after extraction.
- Email header transmission dates are explicitly demarcated as message delivery metadata
  to prevent misinterpreting email send timestamps as complaint incident dates.
"""

import io
import os
import email
from email import policy
from email.parser import BytesParser
from typing import Tuple

from pypdf import PdfReader
import docx

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".eml"}


class DocumentExtractionError(Exception):
    """Raised when document parsing fails or file contains unreadable content."""
    pass


class UnsupportedFileTypeError(DocumentExtractionError):
    """Raised when file extension is not in SUPPORTED_EXTENSIONS."""
    pass


class FileSizeLimitExceededError(DocumentExtractionError):
    """Raised when file size exceeds MAX_FILE_SIZE_BYTES."""
    pass


def validate_file_metadata(filename: str, file_size: int) -> str:
    """
    Validates file extension and size constraints.
    Returns lowercase file extension with leading dot (e.g. '.pdf').
    """
    if not filename:
        raise UnsupportedFileTypeError("Filename cannot be empty.")

    _, ext = os.path.splitext(filename)
    ext_clean = ext.lower().strip()

    if ext_clean not in SUPPORTED_EXTENSIONS:
        supported_str = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext_clean or 'unknown'}'. Supported formats: {supported_str}."
        )

    if file_size > MAX_FILE_SIZE_BYTES:
        raise FileSizeLimitExceededError(
            f"File size ({file_size / (1024 * 1024):.1f} MB) exceeds the maximum 10 MB limit."
        )

    return ext_clean


def extract_pdf_text(file_bytes: bytes) -> str:
    """
    Extracts embedded text streams from PDF files using pypdf.
    If no text can be extracted (e.g. scanned image or blank page),
    raises DocumentExtractionError indicating it may be scanned/image-only.
    """
    try:
        stream = io.BytesIO(file_bytes)
        reader = PdfReader(stream)
        pages_text = []

        for idx, page in enumerate(reader.pages):
            text = page.extract_text()
            if text and text.strip():
                pages_text.append(text.strip())

        full_text = "\n\n".join(pages_text).strip()

        # Check if extracted text is empty or whitespace-only (no arbitrary <5 char rule)
        if not full_text:
            raise DocumentExtractionError(
                "Could not extract readable text from this PDF. The file may be scanned/image-only."
            )

        return full_text
    except DocumentExtractionError:
        raise
    except Exception as e:
        raise DocumentExtractionError(f"Failed to parse PDF document: {str(e)}")


def extract_docx_text(file_bytes: bytes) -> str:
    """
    Extracts text paragraphs and table cell contents from Word .docx files.
    """
    try:
        stream = io.BytesIO(file_bytes)
        doc = docx.Document(stream)
        text_blocks = []

        # Extract standard paragraphs
        for paragraph in doc.paragraphs:
            p_text = paragraph.text.strip()
            if p_text:
                text_blocks.append(p_text)

        # Extract table rows
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Deduplicate adjacent cells caused by merged horizontal cells
                    deduped_cells = []
                    for c in row_cells:
                        if not deduped_cells or deduped_cells[-1] != c:
                            deduped_cells.append(c)
                    text_blocks.append(" | ".join(deduped_cells))

        full_text = "\n".join(text_blocks).strip()

        if not full_text:
            raise DocumentExtractionError("The uploaded DOCX document contains no readable text.")

        return full_text
    except DocumentExtractionError:
        raise
    except Exception as e:
        raise DocumentExtractionError(f"Failed to parse DOCX document: {str(e)}")


def extract_txt_text(file_bytes: bytes) -> str:
    """
    Extracts plain text with UTF-8 decoding, falling back to Latin-1.
    """
    try:
        try:
            full_text = file_bytes.decode("utf-8").strip()
        except UnicodeDecodeError:
            full_text = file_bytes.decode("latin-1", errors="replace").strip()

        if not full_text:
            raise DocumentExtractionError("The uploaded text document is empty.")

        return full_text
    except DocumentExtractionError:
        raise
    except Exception as e:
        raise DocumentExtractionError(f"Failed to decode text document: {str(e)}")


def extract_eml_text(file_bytes: bytes) -> str:
    """
    Extracts email transmission metadata and body text from RFC 822 .eml messages.
    Headers are clearly labeled as email metadata so the LLM does not confuse email
    send timestamps with the pharmaceutical complaint observation/received date.
    """
    try:
        msg = BytesParser(policy=policy.default).parsebytes(file_bytes)

        headers = []
        subject = msg.get("Subject", "").strip()
        from_hdr = msg.get("From", "").strip()
        to_hdr = msg.get("To", "").strip()
        date_hdr = msg.get("Date", "").strip()

        if from_hdr:
            headers.append(f"Email Sender (From): {from_hdr}")
        if to_hdr:
            headers.append(f"Email Recipient (To): {to_hdr}")
        if subject:
            headers.append(f"Email Subject: {subject}")
        if date_hdr:
            # Explicitly annotate that this is email transmission metadata
            headers.append(
                f"Email Message Sent Timestamp (Metadata only, not complaint date unless explicitly confirmed in text): {date_hdr}"
            )

        body_text = ""
        # Prefer plain text body, fallback to HTML
        body_part = msg.get_body(preferencelist=("plain", "html"))
        if body_part:
            body_text = body_part.get_content().strip()
        else:
            # Fallback walk across multipart boundaries
            for part in msg.walk():
                if part.get_content_type() == "text/plain":
                    content = part.get_content()
                    if content:
                        body_text = content.strip()
                        break

        sections = []
        if headers:
            sections.append("[Email Transmission Metadata]\n" + "\n".join(headers))
        if body_text:
            sections.append("[Email Message Body]\n" + body_text)

        full_text = "\n\n".join(sections).strip()

        if not full_text or not body_text:
            raise DocumentExtractionError("The uploaded EML document contains no readable email body content.")

        return full_text
    except DocumentExtractionError:
        raise
    except Exception as e:
        raise DocumentExtractionError(f"Failed to parse EML document: {str(e)}")


def extract_text_from_document(filename: str, file_bytes: bytes) -> Tuple[str, str]:
    """
    Master dispatcher for document text extraction.
    Validates file extension and size, then dispatches to format-specific parser.
    Returns tuple: (extracted_text, file_extension).
    """
    file_size = len(file_bytes)
    ext = validate_file_metadata(filename, file_size)

    if ext == ".pdf":
        extracted_text = extract_pdf_text(file_bytes)
    elif ext == ".docx":
        extracted_text = extract_docx_text(file_bytes)
    elif ext == ".txt":
        extracted_text = extract_txt_text(file_bytes)
    elif ext == ".eml":
        extracted_text = extract_eml_text(file_bytes)
    else:
        raise UnsupportedFileTypeError(f"Unsupported file format '{ext}'.")

    return extracted_text, ext
